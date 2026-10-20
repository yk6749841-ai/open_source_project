"""Excel export — one row per commitment number.

A form with several commitment numbers becomes several rows (the project
requirement, 16/10/2026): every row repeats the form's other fields, and only
the commitment number changes. Several central numbers stay together in one
cell, so rows never multiply into a cross product.

Everything is written as text. Identifiers keep their leading zeros, dates are
not re-parsed into Excel serial numbers, and a value starting with a formula
character can never execute when the lab opens the file.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.fields import FIELDS, MULTI_SEPARATOR
from app.models import STATE_REVIEW, Batch, FormRecord
from app.validators import REASON_TEXT

# נוסחאות: ערך שמתחיל באחד מאלה ייכתב כטקסט מנוטרל ולא ירוץ כנוסחה.
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

NAVY = "1B3A5C"
AMBER_FILL = "FFF1DC"
AMBER_TEXT = "8A4B10"
LINE = "D8E0E8"

EXTRA_COLUMNS = ("קובץ מקור", "עמוד בקובץ", "שדות לבדיקה", "סיבות")


def neutralize(value: str) -> str:
    """Stop a spreadsheet from treating a scanned value as a formula."""
    if value and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


COMMITMENT_KEY = "commitment_numbers"


def record_rows(record: FormRecord) -> list[list[str]]:
    """One row per commitment number; one row when there are none or one."""
    base = record_row(record)
    commitments = record.fields[COMMITMENT_KEY].values
    if len(commitments) < 2:
        return [base]
    column = next(i for i, f in enumerate(FIELDS) if f.key == COMMITMENT_KEY)
    return [[*base[:column], number, *base[column + 1:]] for number in commitments]


def record_row(record: FormRecord) -> list[str]:
    cells = [MULTI_SEPARATOR.join(record.fields[f.key].values) for f in FIELDS]
    flagged = [f for f in record.fields.values() if f.state == STATE_REVIEW]
    reasons = sorted({REASON_TEXT.get(r, r) for f in flagged for r in f.reasons})
    return [
        *cells,
        record.source_name,
        str(record.page_number),
        str(len(flagged)),
        "; ".join(reasons),
    ]


def headers() -> list[str]:
    """The header row of the file — the same list for the file and for the preview."""
    return [f.label_he for f in FIELDS] + list(EXTRA_COLUMNS)


def preview(batch: Batch) -> dict:
    """Exactly the data rows the file will contain, for the screen to show before download.

    Built from the same `record_rows` and `neutralize` as the workbook, so what the
    secretary sees is what the lab gets.
    """
    rows = []
    for number, record in enumerate(batch.records, start=1):
        for values in record_rows(record):
            rows.append({
                "form": number,
                "record_id": record.record_id,
                "cells": [neutralize(v) for v in values],
            })
    return {
        "headers": headers(),
        "rows": rows,
        "forms": len(batch.records),
        "row_count": len(rows),
        "sources": len(batch.accepted_sources),
        "commitment_column": next(i for i, f in enumerate(FIELDS) if f.key == COMMITMENT_KEY),
    }


def build_workbook(batch: Batch) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "התחייבויות"
    ws.sheet_view.rightToLeft = True

    widths = [14, 16, 16, 18, 20, 22, 20, 30, 12, 12, 46]
    thin = Side(style="thin", color=LINE)

    for index, (label, width) in enumerate(zip(headers(), widths), start=1):
        cell = ws.cell(row=1, column=index, value=label)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(index)].width = width
        ws.column_dimensions[get_column_letter(index)].number_format = "@"
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    row_index = 1
    for record in batch.records:
        for values in record_rows(record):
            row_index += 1
            for column_index, value in enumerate(values, start=1):
                cell = ws.cell(row=row_index, column=column_index, value=neutralize(value))
                cell.number_format = "@"  # טקסט: אפסים מובילים ותאריכים נשמרים כפי שהם
                cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
                cell.border = Border(bottom=thin)

                spec = FIELDS[column_index - 1] if column_index <= len(FIELDS) else None
                if spec and record.fields[spec.key].state == STATE_REVIEW:
                    # כתום = נדרשת בדיקה. אותה משמעות כמו במסך — בכל שורות הטופס.
                    cell.fill = PatternFill("solid", fgColor=AMBER_FILL)
                    cell.font = Font(color=AMBER_TEXT)

    _add_notice(ws, row_index + 2, batch, rows=row_index - 1)
    return wb


def _add_notice(ws, row: int, batch: Batch, *, rows: int) -> None:
    sources = len(batch.accepted_sources)
    lines = [
        "דמו להאקתון — מיועד לטפסי הדוגמה בלבד. אין להסתמך על הקובץ ללא בדיקה.",
        f"מקור: {batch.display_name} · {sources} קבצים · {len(batch.records)} טפסים · "
        f"{rows} שורות (שורה לכל מספר התחייבות) · "
        f"{batch.review_total} שדות מסומנים לבדיקה",
        "תאים כתומים דורשים בדיקה אנושית, כולל שדות ריקים. אין כאן טענת נכונות.",
    ]
    if batch.rejected_sources:
        names = ", ".join(f"{s.filename} ({s.error})" for s in batch.rejected_sources)
        lines.append(f"קבצים שלא עובדו: {names}")
    for offset, text in enumerate(lines):
        cell = ws.cell(row=row + offset, column=1, value=text)
        cell.font = Font(size=10, color="5A6B7B", italic=offset > 0)
        cell.alignment = Alignment(horizontal="right")


def workbook_bytes(batch: Batch) -> bytes:
    buffer = BytesIO()
    build_workbook(batch).save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------ ייצוא "רשומות במערכת"
#
# רק טפסים מאושרים, באותה צורה כמו הקובץ הרגיל: שורה לכל מספר התחייבות. השורות
# מגיעות מוכנות ממסד הנתונים (db.rows_for_export).

SAVED_KEYS = ("first_name", "last_name", "national_id", "central_numbers",
              "sample_received_date", "sending_entity", "commitment_number")
SAVED_HEADERS = [f.label_he for f in FIELDS] + ["קובץ מקור", "עמוד בקובץ", "אושר"]


def saved_cells(row: dict) -> list[str]:
    return [str(row.get(k) or "") for k in SAVED_KEYS] + [
        str(row.get("source_name") or ""), str(row.get("page_number") or ""), str(row.get("approved_local") or ""),
    ]


def saved_workbook_bytes(rows: list[dict], title: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "טפסים מאושרים"
    ws.sheet_view.rightToLeft = True
    thin = Side(style="thin", color=LINE)

    widths = [14, 16, 16, 18, 20, 22, 20, 30, 12, 18]
    for index, (label, width) in enumerate(zip(SAVED_HEADERS, widths), start=1):
        cell = ws.cell(row=1, column=index, value=label)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    for row_index, row in enumerate(rows, start=2):
        for column_index, value in enumerate(saved_cells(row), start=1):
            cell = ws.cell(row=row_index, column=column_index, value=neutralize(value))
            cell.number_format = "@"  # טקסט: אפסים מובילים ותאריכים נשמרים כפי שהם
            cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)

    note = ws.cell(row=len(rows) + 3, column=1, value=title)
    note.font = Font(size=10, color="5A6B7B", italic=True)
    note.alignment = Alignment(horizontal="right")
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def saved_csv_bytes(rows: list[dict]) -> bytes:
    """UTF-8 with BOM, so Excel opens the Hebrew correctly. Every value neutralised."""
    import csv
    import io

    text = io.StringIO()
    writer = csv.writer(text)
    writer.writerow(SAVED_HEADERS)
    for row in rows:
        writer.writerow([neutralize(v) for v in saved_cells(row)])
    return ("﻿" + text.getvalue()).encode("utf-8")


def saved_json(rows: list[dict]) -> list[dict]:
    """The same rows, as objects with English keys."""
    return [
        {**{k: row.get(k) or "" for k in SAVED_KEYS},
         "source_name": row.get("source_name") or "", "page_number": row.get("page_number"),
         "approved_at": row.get("approved_local") or ""}
        for row in rows
    ]
