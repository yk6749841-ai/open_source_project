"""The "רשומות במערכת" panel: only approved forms, one row per commitment number, by day."""

from __future__ import annotations

import io
import json
import sqlite3
from datetime import datetime, timedelta, timezone

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from app import db, pipeline, store
from app.models import STATE_REVIEW, Batch, FormRecord, SourceFile

TODAY = datetime.now(timezone.utc).replace(hour=0, minute=0, second=0, microsecond=0)
RANGE = {"start": TODAY.isoformat(), "end": (TODAY + timedelta(days=1)).isoformat()}


@pytest.fixture(autouse=True)
def _isolated(tmp_path, monkeypatch):
    monkeypatch.setattr("app.store.BATCHES_DIR", tmp_path / "batches")
    monkeypatch.setattr("app.config.BATCHES_DIR", tmp_path / "batches")
    monkeypatch.setattr("app.gold.save", lambda *a, **k: None)  # לא נוגעים בזהב


def _batch(batch_id: str, *forms: dict, shade: int = 0) -> Batch:
    batch = Batch(batch_id=batch_id, status="completed")
    source = SourceFile(index=1, filename=f"{batch_id}.pdf", kind="pdf", page_count=len(forms))
    batch.sources.append(source)
    for number, values in enumerate(forms, start=1):
        page = store.source_dir(batch_id, 1) / "pages" / f"page-{number:03d}.png"
        page.parent.mkdir(parents=True, exist_ok=True)
        # כל עמוד שונה בפיקסלים — כמו טפסים אמיתיים (אותו עמוד = אותו טופס)
        Image.new("RGB", (60, 80), (255, 255 - number, 255 - shade)).save(page)
        record = FormRecord.empty(batch_id, source, number)
        for key, vals in values.items():
            record.fields[key].values = list(vals)
        batch.records.append(record)
    store.put(batch)
    return batch


def _client() -> TestClient:
    from app.api import app
    return TestClient(app)


def _approve(record: FormRecord) -> None:
    response = _client().post(f"/api/records/{record.record_id}/verify", json={"version": record.version})
    assert response.status_code == 200, response.text


def test_reading_and_fixing_save_nothing():
    batch = _batch("bat_a", {"first_name": ["שרה"]})
    _client().patch(f"/api/records/{batch.records[0].record_id}",
                    json={"field": "first_name", "values": ["שרה"]})
    assert db.list_rows(**RANGE)["rows"] == []


def test_approval_saves_one_row_per_commitment_number():
    batch = _batch("bat_b", {"first_name": ["שרה"], "national_id": ["012345678"],
                             "central_numbers": ["948023", "948024"], "commitment_numbers": ["0111", "0222"]})
    _approve(batch.records[0])

    found = db.list_rows(**RANGE)

    assert [r["commitment_number"] for r in found["rows"]] == ["0111", "0222"]  # אפס מוביל נשמר
    assert {r["central_numbers"] for r in found["rows"]} == {"948023 | 948024"}  # מרכזיים יחד
    assert found["forms"] == 1 and found["total_rows"] == 2
    assert all(r["national_id"] == "012345678" for r in found["rows"])


def test_form_without_commitment_number_is_one_row():
    batch = _batch("bat_c", {"first_name": ["דוד"]})
    _approve(batch.records[0])
    rows = db.list_rows(**RANGE)["rows"]
    assert len(rows) == 1 and rows[0]["commitment_number"] == ""


def test_approving_again_or_uploading_again_never_adds_rows():
    first = _batch("bat_d", {"first_name": ["שרה"], "commitment_numbers": ["1", "2"]})
    _approve(first.records[0])
    _approve(first.records[0])  # אישור חוזר
    again = _batch("bat_e", {"first_name": ["שרה"], "commitment_numbers": ["1", "2"]})  # אותו עמוד, העלאה חדשה
    _approve(again.records[0])

    assert db.list_rows(**RANGE)["total_rows"] == 2


def test_approved_time_is_the_latest_approval(monkeypatch):
    batch = _batch("bat_t", {"first_name": ["שרה"], "commitment_numbers": ["1"]})
    monkeypatch.setattr("app.db.now", lambda: "2026-10-20T08:00:00+00:00")
    _approve(batch.records[0])
    monkeypatch.setattr("app.db.now", lambda: "2026-10-20T09:30:00+00:00")
    _approve(batch.records[0])
    assert [r["approved_at"] for r in db.list_rows(**RANGE)["rows"]] == ["2026-10-20T09:30:00+00:00"]


def test_a_fix_after_approval_replaces_the_rows():
    batch = _batch("bat_f", {"first_name": ["שרה"], "commitment_numbers": ["1", "2"]})
    record = batch.records[0]
    _approve(record)
    record.fields["commitment_numbers"].values = ["3"]
    _approve(record)
    assert [r["commitment_number"] for r in db.list_rows(**RANGE)["rows"]] == ["3"]


def test_the_panel_loads_only_the_chosen_days():
    batch = _batch("bat_g", {"first_name": ["היום"]})
    _approve(batch.records[0])
    yesterday = {"start": (TODAY - timedelta(days=1)).isoformat(), "end": TODAY.isoformat()}
    week = {"start": (TODAY - timedelta(days=7)).isoformat(), "end": RANGE["end"]}

    assert len(db.list_rows(**RANGE)["rows"]) == 1
    assert db.list_rows(**yesterday)["rows"] == []
    assert len(db.list_rows(**week)["rows"]) == 1
    assert _client().get("/api/saved").status_code == 400  # בלי טווח — אין שליפה של כל המאגר


def test_filters_search_sender_and_received_date():
    batch = _batch("bat_h",
                   {"first_name": ["רחל"], "sending_entity": ["מכבי"], "sample_received_date": ["05/10/2026"]},
                   {"first_name": ["לאה"], "sending_entity": ["כללית"], "sample_received_date": ["18/10/2026"]})
    for record in batch.records:
        _approve(record)

    names = lambda **f: [r["first_name"] for r in db.list_rows(**RANGE, **f)["rows"]]  # noqa: E731
    assert names(q="רחל") == ["רחל"]
    assert names(sender="כללית") == ["לאה"]
    assert names(received_from="2026-10-10") == ["לאה"]
    assert names(received_to="2026-10-10") == ["רחל"]
    assert db.list_rows(**RANGE)["senders"] == ["כללית", "מכבי"]


def test_same_commitment_number_in_another_form_is_flagged_not_merged():
    one = _batch("bat_i", {"first_name": ["א"], "commitment_numbers": ["7254647"]})
    two = _batch("bat_j", {"first_name": ["ב"], "commitment_numbers": ["7254647"]},
                 {"first_name": ["ג"], "commitment_numbers": ["111"]}, shade=9)
    for record in (*one.records, *two.records):
        _approve(record)

    rows = {r["first_name"]: r for r in db.list_rows(**RANGE)["rows"]}
    assert len(rows) == 3  # אף אחד לא נמחק
    assert [d["source_name"] for d in rows["ב"]["duplicates"]] == ["bat_i.pdf"]
    assert rows["ג"]["duplicates"] == []


def test_export_is_the_same_rows_as_text_with_formulas_neutralised():
    batch = _batch("bat_k", {"first_name": ["=cmd"], "national_id": ["012345678"], "commitment_numbers": ["1", "2"]})
    _approve(batch.records[0])
    client = _client()

    sheet = load_workbook(io.BytesIO(client.get("/api/saved/export", params={**RANGE, "format": "xlsx"}).content)).active
    rows = [[c.value for c in row] for row in sheet.iter_rows(min_row=2, max_row=3)]
    assert [r[0] for r in rows] == ["'=cmd", "'=cmd"] and [r[6] for r in rows] == ["1", "2"]
    assert rows[0][2] == "012345678"

    csv_text = client.get("/api/saved/export", params={**RANGE, "format": "csv"}).content.decode("utf-8-sig")
    assert "'=cmd" in csv_text and "012345678" in csv_text
    data = json.loads(client.get("/api/saved/export", params={**RANGE, "format": "json"}).content)
    assert [d["commitment_number"] for d in data] == ["1", "2"]


def test_reopen_after_restart_without_reading_again():
    batch = _batch("bat_l", {"first_name": ["שרה"]}, {"first_name": ["דוד"]})
    for record in batch.records:
        _approve(record)
    wanted = batch.records[1]
    key = next(r["form_key"] for r in db.list_rows(**RANGE)["rows"] if r["first_name"] == "דוד")
    store.delete("bat_l")  # כמו הפעלה מחדש: האצווה והקבצים שלה נעלמו

    opened = _client().post(f"/api/saved/{key}/open")

    assert opened.status_code == 200
    body = opened.json()
    assert body["focus"] == wanted.record_id and len(body["records"]) == 2
    assert pipeline.page_image_path("bat_l", 1, 2).is_file()  # העמוד הוחזר מהעותק השמור


def test_delete_removes_the_rows_and_every_image():
    batch = _batch("bat_m", {"first_name": ["שרה"]}, {"first_name": ["דוד"]})
    for record in batch.records:
        _approve(record)
    key = next(r["form_key"] for r in db.list_rows(**RANGE)["rows"] if r["first_name"] == "שרה")
    page = pipeline.page_image_path("bat_m", 1, 1)
    kept = db.PAGES_DIR / f"{key}.png"
    assert kept.is_file()

    assert _client().delete(f"/api/saved/{key}").status_code == 200

    assert [r["first_name"] for r in db.list_rows(**RANGE)["rows"]] == ["דוד"]
    assert not page.exists() and not kept.exists()
    assert key not in [r["form_key"] for r in db.list_approved()]
    assert pipeline.page_image_path("bat_m", 1, 2).is_file()  # הטופס השני לא נפגע


def test_forms_approved_before_the_panel_appear_in_it(tmp_path):
    path = tmp_path / "old.db"
    stamp = datetime.now(timezone.utc).isoformat(timespec="seconds")
    with sqlite3.connect(path) as conn:
        conn.execute(db._SCHEMA)
        conn.execute("CREATE TABLE records (x)")  # מה שהגרסה הקודמת השאירה — יימחק
        conn.execute(
            "INSERT INTO approved_forms (form_key, first_name, central_numbers, commitment_numbers, "
            "source_file, page_number, approved_at, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
            ("old1", "ותיק", '["948023"]', '["5", "6"]', "ישן.pdf", 1, stamp, stamp))

    found = db.list_rows(**RANGE, path=path)

    assert [r["commitment_number"] for r in found["rows"]] == ["5", "6"]
    assert all(not r["can_open"] for r in found["rows"])  # אפשר לראות ולייצא, לא לפתוח
    with sqlite3.connect(path) as conn:
        assert not conn.execute("SELECT name FROM sqlite_master WHERE name = 'records'").fetchall()


def test_review_screen_shows_no_saved_data():
    batch = _batch("bat_n", {"first_name": ["שרה"]})
    batch.records[0].fields["first_name"].state = STATE_REVIEW
    body = _client().get("/api/batches/bat_n").json()
    assert "duplicates" not in body["records"][0]
